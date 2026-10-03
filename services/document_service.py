from pathlib import Path
from pypdf import PdfReader
from docx import Document

ALLOWED_EXTENSIONS = {
    "pdf",
    "docx",
    "jpg",
    "jpeg",
    "png"
}

IMAGE_EXTENSIONS = {
    "jpg",
    "jpeg",
    "png"
}


def allowed_file(filename: str) -> bool:
    """Checks whether the file extension is supported."""
    if not filename:
        return False
    extension = Path(filename).suffix.lower().replace(".", "")
    return extension in ALLOWED_EXTENSIONS


def is_image_file(filename: str) -> bool:
    """Checks if the file is an image."""
    if not filename:
        return False
    extension = Path(filename).suffix.lower().replace(".", "")
    return extension in IMAGE_EXTENSIONS


def extract_pdf_text(filepath: str) -> str:
    """Extracts text content from a PDF file using pypdf."""
    reader = PdfReader(filepath)
    pages = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages.append(text)
    return "\n".join(pages).strip()


def extract_docx_text(filepath: str) -> str:
    """Extracts text content from a DOCX file, including tables."""
    doc = Document(filepath)
    content_blocks = []

    # Extract text from paragraphs
    for paragraph in doc.paragraphs:
        txt = paragraph.text.strip()
        if txt:
            content_blocks.append(txt)

    # Extract text from tables (common for invoices and receipts)
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                # Remove consecutive duplicates in merged cells
                cleaned_cells = []
                for cell_text in row_cells:
                    if not cleaned_cells or cleaned_cells[-1] != cell_text:
                        cleaned_cells.append(cell_text)
                content_blocks.append(" | ".join(cleaned_cells))

    return "\n".join(content_blocks).strip()


def extract_document_text(filepath: str) -> str:
    """
    Extracts text depending on document type.
    For PDF: uses pypdf.
    For DOCX: uses python-docx.
    For images: returns empty string (keeps system ready for future OCR/Vision without inventing text).
    """
    ext = Path(filepath).suffix.lower()

    if ext == ".pdf":
        return extract_pdf_text(filepath)
    elif ext == ".docx":
        return extract_docx_text(filepath)
    elif ext.replace(".", "") in IMAGE_EXTENSIONS:
        # Per requirement: Keep upload system ready for future OCR/Vision. Do not invent extracted text from image.
        return ""

    raise ValueError(f"Unsupported file type: {ext}")